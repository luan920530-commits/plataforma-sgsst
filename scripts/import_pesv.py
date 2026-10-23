#!/usr/bin/env python3
"""Importa el paquete PESV a recursos seguros y autocontenidos, sin ejecutar OOXML."""
from __future__ import annotations

import collections
import csv
import hashlib
import html
import io
import json
import re
import shutil
import tempfile
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path, PurePosixPath

import py7zr
from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / "pesvial.7z"
OUTPUT = ROOT / "resources" / "pesv-library"
PHASES = {"1": "Planear", "2": "Hacer", "3": "Verificar", "4": "Actuar"}
ROUTES = {
    1: "/pesv?paso=1", 2: "/pesv?paso=2", 3: "/pesv?paso=3", 4: "/pesv?paso=4",
    5: "/pesv?paso=5", 6: "/matriz-ipevr", 7: "/indicadores", 8: "/programas-prevencion",
    9: "/plan-anual", 10: "/capacitaciones", 11: "/equipo", 12: "/emergencias",
    13: "/autoreportes", 14: "/pesv?paso=14", 15: "/pesv?paso=15", 16: "/pesv?paso=16",
    17: "/pesv?paso=17", 18: "/cumplimiento-documental", 19: "/cumplimiento-documental", 20: "/indicadores",
    21: "/autoreportes", 22: "/cumplimiento-documental", 23: "/hallazgos", 24: "/cumplimiento-documental",
}
MACRO_TOKENS = ("vbaproject", "vbadata", "activex", "embeddings", "oleobject", "customui", "macrosheet")
RISK_FORMULA = re.compile(r"\b(WEBSERVICE|RTD|DDE|DDEAUTO|CALL|REGISTER\.ID|SHELL|EXEC)\s*\(|\[[0-9]+\]", re.I)
STEP_PATTERN = re.compile(r"PASO\s+(\d+)\s+(.+)", re.I)


def step_info(relative: str) -> tuple[int, str, str]:
    parts = PurePosixPath(relative).parts
    match = None
    for part in reversed(parts[1:-1]):
        match = STEP_PATTERN.search(part)
        if match:
            break
    if not match:
        raise RuntimeError(f"No se pudo identificar el paso PESV: {relative}")
    number = int(match.group(1))
    if number < 1 or number > 24:
        raise RuntimeError(f"Paso fuera de rango: {relative}")
    phase = PHASES[parts[0].split(".", 1)[0]]
    return number, re.sub(r"\s+", " ", match.group(2)).strip(), phase


def security_scan(path: Path, relative: str) -> dict:
    result = {"macro_parts": [], "external_parts": [], "external_relationships": 0,
              "external_targets": [], "risk_formulas": 0, "defined_name_external_references": 0, "hyperlinks": 0}
    with zipfile.ZipFile(path) as package:
        for name in package.namelist():
            low = name.lower()
            if any(token in low for token in MACRO_TOKENS):
                result["macro_parts"].append(name)
            if "externallinks" in low:
                result["external_parts"].append(name)
            if not low.endswith((".xml", ".rels")):
                continue
            data = package.read(name)
            if low.endswith(".rels"):
                import xml.etree.ElementTree as ET
                try:
                    root = ET.fromstring(data)
                except ET.ParseError:
                    continue
                for node in root:
                    if node.attrib.get("TargetMode", "").lower() == "external":
                        result["external_relationships"] += 1
                        result["external_targets"].append(node.attrib.get("Target", ""))
                        if node.attrib.get("Type", "").endswith("/hyperlink"):
                            result["hyperlinks"] += 1
            if low.startswith("xl/worksheets/") and low.endswith(".xml"):
                result["risk_formulas"] += len(RISK_FORMULA.findall(data.decode("utf-8", "replace")))
            elif low == "xl/workbook.xml":
                result["defined_name_external_references"] += len(RISK_FORMULA.findall(data.decode("utf-8", "replace")))
            elif low == "word/document.xml":
                result["risk_formulas"] += len(RISK_FORMULA.findall(data.decode("utf-8", "replace")))
    return result


def scalar(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, timedelta):
        return value.total_seconds()
    return str(value)


def formula_text(cell) -> str:
    value = cell.value
    if isinstance(value, str):
        return value[1:] if value.startswith("=") else value
    if hasattr(value, "text"):
        return str(value.text).lstrip("=")
    return str(value).lstrip("=")


def style(cell) -> dict:
    if not cell.has_style:
        return {}
    result = {}
    if cell.font.bold:
        result["b"] = 1
    if cell.font.italic:
        result["i"] = 1
    border = any(getattr(cell.border, side).style for side in ("left", "right", "top", "bottom"))
    if border:
        result["border"] = 1
    if cell.fill.fill_type:
        rgb = cell.fill.fgColor.rgb
        if isinstance(rgb, str):
            value = rgb[-6:]
            if re.fullmatch(r"[A-Fa-f0-9]{6}", value) and value.upper() != "000000":
                result["fill"] = value.upper()
    if cell.alignment.horizontal in ("left", "center", "right", "justify"):
        result["align"] = cell.alignment.horizontal
    if cell.protection.locked is False:
        result["unlocked"] = 1
    return result


def meaningful(cell) -> bool:
    return cell.value is not None or style(cell) != {}


def validation_options(workbook, worksheet, validation) -> list[str]:
    if validation.type != "list" or not validation.formula1:
        return []
    source = validation.formula1.strip()
    if source.startswith('"') and source.endswith('"'):
        return [item.strip() for item in next(csv.reader([source[1:-1]])) if item.strip()]
    if "[" in source or "]" in source:
        return []
    try:
        if "!" in source:
            sheet_name, reference = source.rsplit("!", 1)
            sheet_name = sheet_name.strip("'").replace("''", "'")
            target = workbook[sheet_name]
        else:
            target, reference = worksheet, source
        min_col, min_row, max_col, max_row = range_boundaries(reference.replace("$", ""))
        if max_col - min_col > 100 or max_row - min_row > 1000:
            return []
        values = []
        for row in target.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if cell.value is not None and str(cell.value) not in values:
                    values.append(str(cell.value))
        return values
    except (ValueError, KeyError, TypeError):
        return []


def convert_xlsx(path: Path) -> tuple[dict, dict, list[str]]:
    # No se conservan ni resuelven vínculos externos. Las cachés externalLink nunca se leen.
    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, read_only=False, keep_links=False)
    output = {"sheets": []}
    stats = {"sheets": 0, "cells": 0, "formulas": 0, "images": 0, "charts": 0, "validations": 0,
             "residual_format_cells_excluded": 0}
    warnings = []
    for worksheet in workbook.worksheets:
        cached_sheet = cached[worksheet.title]
        value_cells = [cell for cell in worksheet._cells.values() if cell.value is not None]
        sheet = {"name": worksheet.title, "state": worksheet.sheet_state,
                 "rows": [], "columns": [], "cells": {}, "merged": {},
                 "images": len(worksheet._images), "charts": len(worksheet._charts)}
        stats["sheets"] += 1
        stats["images"] += len(worksheet._images)
        stats["charts"] += len(worksheet._charts)
        if not value_cells:
            output["sheets"].append(sheet)
            continue
        # Todo valor/fórmula se conserva, sin límite de coordenada. Solo el formato
        # vacío residual queda fuera del área útil delimitada por contenido y combinaciones.
        min_row = min(cell.row for cell in value_cells)
        max_row = max(cell.row for cell in value_cells)
        min_col = min(cell.column for cell in value_cells)
        max_col = max(cell.column for cell in value_cells)
        merges = []
        for merged in worksheet.merged_cells.ranges:
            if merged.max_row < min_row or merged.min_row > max_row or merged.max_col < min_col or merged.min_col > max_col:
                continue
            merges.append(merged)
            min_row, max_row = min(min_row, merged.min_row), max(max_row, merged.max_row)
            min_col, max_col = min(min_col, merged.min_col), max(max_col, merged.max_col)
        rows_with_content, cols_with_content = set(), set()
        for cell in list(worksheet._cells.values()):
            has_value = cell.value is not None
            if not has_value and not (min_row <= cell.row <= max_row and min_col <= cell.column <= max_col):
                if cell.has_style:
                    stats["residual_format_cells_excluded"] += 1
                continue
            if not meaningful(cell):
                continue
            coordinate = cell.coordinate
            entry = {}
            if cell.data_type == "f":
                formula = formula_text(cell)
                if RISK_FORMULA.search(formula):
                    warnings.append(f"Se excluyó una fórmula externa o no segura en {worksheet.title}!{coordinate}.")
                    entry["v"] = ""
                else:
                    entry["f"] = formula
                    entry["v"] = scalar(cached_sheet[coordinate].value) or ""
                    stats["formulas"] += 1
            elif cell.value is not None:
                entry["v"] = scalar(cell.value)
                stats["cells"] += 1
            cell_style = style(cell)
            if cell_style:
                entry["s"] = cell_style
            sheet["cells"][coordinate] = entry
            rows_with_content.add(cell.row)
            cols_with_content.add(cell.column)
        for merged in merges:
            anchor = f"{get_column_letter(merged.min_col)}{merged.min_row}"
            if anchor not in sheet["cells"]:
                sheet["cells"][anchor] = {"s": {}}
            sheet["cells"][anchor]["colspan"] = merged.max_col - merged.min_col + 1
            sheet["cells"][anchor]["rowspan"] = merged.max_row - merged.min_row + 1
            rows_with_content.add(merged.min_row)
            cols_with_content.add(merged.min_col)
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    if row == merged.min_row and col == merged.min_col:
                        continue
                    sheet["merged"][f"{get_column_letter(col)}{row}"] = anchor
        validations = getattr(worksheet.data_validations, "dataValidation", []) if worksheet.data_validations else []
        for validation in validations:
            options = validation_options(workbook, worksheet, validation)
            if not options:
                continue
            stats["validations"] += 1
            for cell_range in validation.sqref.ranges:
                for row in range(cell_range.min_row, cell_range.max_row + 1):
                    for col in range(cell_range.min_col, cell_range.max_col + 1):
                        coordinate = f"{get_column_letter(col)}{row}"
                        if coordinate in sheet["cells"] and "f" not in sheet["cells"][coordinate]:
                            sheet["cells"][coordinate]["options"] = options
        sheet["rows"] = sorted(rows_with_content)
        sheet["columns"] = sorted(cols_with_content)
        output["sheets"].append(sheet)
    workbook.close()
    cached.close()
    if stats["images"] or stats["charts"]:
        warnings.append("Las imágenes o gráficos no se reproducen en el editor; revisa el original antes de aprobar.")
    if stats["residual_format_cells_excluded"]:
        warnings.append(f"Se excluyeron {stats['residual_format_cells_excluded']} celdas vacías de formato residual fuera del área útil; todos los valores y fórmulas se conservaron.")
    return output, stats, list(dict.fromkeys(warnings))


def paragraph_html(paragraph: Paragraph) -> str:
    if not paragraph.text.strip():
        return ""
    content = ""
    for run in paragraph.runs:
        value = html.escape(run.text).replace("\n", "<br>")
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        content += value
    if not content:
        content = html.escape(paragraph.text)
    style_name = (paragraph.style.name if paragraph.style else "").lower()
    tag = "h2" if "heading 1" in style_name or "título 1" in style_name else ("h3" if "heading" in style_name or "título" in style_name else "p")
    return f"<{tag}>{content}</{tag}>"


def blocks(parent):
    if isinstance(parent, DocumentType):
        element = parent.element.body
    elif isinstance(parent, _Cell):
        element = parent._tc
    else:
        raise TypeError(type(parent))
    for child in element.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def table_html(table: Table) -> tuple[str, dict]:
    """Conserva bloques, runs y tablas hijas sin convertir la celda a texto plano."""
    output = ["<table><tbody>"]
    stats = {"paragraphs": 0, "tables": 1, "characters": 0}
    for row in table.rows:
        output.append("<tr>")
        for cell in row.cells:
            output.append("<td>")
            for block in blocks(cell):
                if isinstance(block, Paragraph):
                    rendered = paragraph_html(block)
                    if rendered:
                        output.append(rendered)
                        stats["paragraphs"] += 1
                        stats["characters"] += len(block.text)
                else:
                    rendered, nested = table_html(block)
                    output.append(rendered)
                    for key in stats:
                        stats[key] += nested[key]
            output.append("</td>")
        output.append("</tr>")
    output.append("</tbody></table>")
    return "".join(output), stats


def convert_docx(path: Path, title: str) -> tuple[str, dict, list[str]]:
    document = Document(path)
    body = [f'<article class="sst-document library-import"><header><h1>{html.escape(title)}</h1>',
            f'<p><strong>Fuente importada:</strong> {html.escape(path.name)}. Borrador de referencia; adapte, revise y apruebe antes de usar.</p></header>']
    paragraphs = tables = characters = 0
    for block in blocks(document):
        if isinstance(block, Paragraph):
            rendered = paragraph_html(block)
            if rendered:
                body.append(rendered)
                paragraphs += 1
                characters += len(block.text)
        else:
            rendered, table_stats = table_html(block)
            body.append(rendered)
            paragraphs += table_stats["paragraphs"]
            tables += table_stats["tables"]
            characters += table_stats["characters"]
    body.append("</article>")
    with zipfile.ZipFile(path) as package:
        images = sum(name.startswith("word/media/") for name in package.namelist())
        charts = sum(name.startswith("word/charts/") and name.endswith(".xml") for name in package.namelist())
    warnings = []
    if images or charts:
        warnings.append("Las imágenes, gráficos o elementos de diseño no se reproducen; revisa el original antes de aprobar.")
    return "".join(body), {"paragraphs": paragraphs, "tables": tables, "characters": characters, "images": images, "charts": charts}, warnings


def main() -> None:
    if not ARCHIVE.is_file():
        raise SystemExit("No se encontró pesvial.7z en la raíz del proyecto")
    with tempfile.TemporaryDirectory(prefix="pesv-import-") as temporary:
        source = Path(temporary)
        with py7zr.SevenZipFile(ARCHIVE) as archive:
            for info in archive.list():
                pure = PurePosixPath(info.filename)
                if pure.is_absolute() or ".." in pure.parts or "\\" in info.filename:
                    raise RuntimeError(f"Ruta insegura en 7z: {info.filename}")
            archive.extractall(source)
        paths = sorted(path for path in source.rglob("*") if path.is_file())
        if len(paths) != 114:
            raise RuntimeError(f"Inventario inesperado: {len(paths)} archivos")
        grouped = {}
        security = {}
        for path in paths:
            relative = path.relative_to(source).as_posix()
            if path.suffix.lower() not in (".xlsx", ".docx"):
                raise RuntimeError(f"Tipo no admitido: {relative}")
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            step_info(relative)
            security[relative] = security_scan(path, relative)
            grouped.setdefault(digest, []).append(path)
        OUTPUT.mkdir(parents=True, exist_ok=True)
        (OUTPUT / "forms").mkdir(exist_ok=True)
        (OUTPUT / "content").mkdir(exist_ok=True)
        steps = {}
        items = []
        all_targets = set()
        for paths_group in grouped.values():
            paths_group.sort(key=lambda p: (step_info(p.relative_to(source).as_posix())[0], p.name.lower()))
            canonical = paths_group[0]
            relative = canonical.relative_to(source).as_posix()
            step, step_name, phase = step_info(relative)
            sources = []
            source_steps = set()
            for path in paths_group:
                rel = path.relative_to(source).as_posix()
                source_step, source_name, source_phase = step_info(rel)
                source_steps.add(source_step)
                steps[source_step] = {"number": source_step, "name": source_name, "phase": source_phase, "route": ROUTES[source_step]}
                sources.append({"path": rel, "step": source_step})
            digest = hashlib.sha256(canonical.read_bytes()).hexdigest()
            identifier = "PESV-" + digest[:12].upper()
            scan = security[relative]
            warnings = []
            if scan["macro_parts"]:
                warnings.append("El original contiene partes activas u OLE; fueron excluidas y nunca se ejecutaron.")
            if scan["external_parts"] or scan["external_relationships"]:
                warnings.append("El original contiene vínculos externos heredados; se excluyeron sus relaciones y cachés. No se resuelven ni actualizan.")
            if scan["risk_formulas"]:
                warnings.append("Se detectaron fórmulas externas o de riesgo; se neutralizaron para edición segura.")
            if scan["defined_name_external_references"]:
                warnings.append("Los nombres definidos vinculados a archivos externos fueron excluidos; no se copiaron sus referencias ni resultados en caché.")
            if canonical.suffix.lower() == ".xlsx":
                content, stats, conversion_warnings = convert_xlsx(canonical)
                kind, extension = "form", "xlsx"
                resource = f"forms/{identifier}.json"
                (OUTPUT / resource).write_text(json.dumps(content, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
                document_type = "Registro"
            else:
                title = canonical.stem.strip()
                content, stats, conversion_warnings = convert_docx(canonical, title)
                kind, extension = "document", "docx"
                resource = f"content/{identifier}.html"
                (OUTPUT / resource).write_text(content, encoding="utf-8")
                document_type = "Procedimiento"
            warnings.extend(conversion_warnings)
            item = {
                "id": identifier, "code": f"PESV-{step:02d}-{digest[:6].upper()}",
                "title": canonical.stem.strip(), "extension": extension, "kind": kind,
                "collection": "pesv", "standards": [], "primary_standard": 0,
                "step": step, "steps": sorted(source_steps), "step_name": step_name,
                "phase": phase, "folder": PurePosixPath(relative).parent.as_posix(),
                "module": f"PESV · Paso {step} · {step_name}", "route": ROUTES[step],
                "content": resource, "sha256": digest, "bytes": canonical.stat().st_size,
                "original_name": canonical.name, "sources": sources,
                "duplicate_entries": len(sources) - 1, "warnings": list(dict.fromkeys(warnings)),
                "stats": stats, "status": "convertido", "document_type": document_type,
            }
            items.append(item)
        items.sort(key=lambda item: (item["step"], item["title"].casefold(), item["id"]))
        for number, item in enumerate(items, 1):
            item["number"] = number
        for rel, scan in security.items():
            all_targets.update(scan["external_targets"])
        manifest = {
            "name": "Biblioteca PESV · 24 pasos", "source_archive": ARCHIVE.name,
            "source_sha256": hashlib.sha256(ARCHIVE.read_bytes()).hexdigest(),
            "source_entries": len(paths), "unique_templates": len(items),
            "duplicate_entries": len(paths) - len(items), "ignored_entries": 0,
            "steps_covered": sorted(steps), "steps": [steps[number] for number in sorted(steps)],
            "counts": dict(collections.Counter(path.suffix.lower().lstrip(".") for path in paths)),
            "security": {
                "macros_activex_ole_files": sum(bool(scan["macro_parts"]) for scan in security.values()),
                "files_with_external_links": sum(bool(scan["external_parts"] or scan["external_relationships"]) for scan in security.values()),
                "external_link_parts_excluded": sum(len(scan["external_parts"]) for scan in security.values()),
                "external_targets_unique_excluded": len(all_targets),
                "risk_formulas_detected": sum(scan["risk_formulas"] for scan in security.values()),
                "external_defined_name_references_excluded": sum(scan["defined_name_external_references"] for scan in security.values()),
                "policy": "Sin ejecución de macros; vínculos externos, nombres definidos externos, cachés y relaciones externas excluidos de los recursos editables.",
            },
            "templates": items,
        }
        if manifest["steps_covered"] != list(range(1, 25)):
            raise RuntimeError("El paquete no cubre los 24 pasos esperados")
        (OUTPUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({key: value for key, value in manifest.items() if key not in ("templates", "steps")}, ensure_ascii=False, indent=2))
        print(f"Conversión completa: {len(items)} plantillas únicas en {OUTPUT}")


if __name__ == "__main__":
    main()
