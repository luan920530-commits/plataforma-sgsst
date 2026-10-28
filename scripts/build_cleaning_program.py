#!/usr/bin/env python3
"""Genera el kit editable PR/FT-LD desde la misma fuente JSON utilizada por PHP.

Dependencias sólo para construir los anexos: python-docx y openpyxl.
No ejecuta macros ni marca actividades como realizadas.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from datetime import date
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exception:
    raise SystemExit(
        "Faltan dependencias para generar DOCX/XLSX. "
        "Instala python-docx y openpyxl en un entorno de construcción."
    ) from exception

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "resources/cleaning-program/program.json"
SEED = ROOT / "resources/extintores-scm/seed.json"
arguments = sys.argv[1:]
if len(arguments) > 1:
    raise SystemExit("Uso: python3 scripts/build_cleaning_program.py [directorio_de_salida]")
OUTPUT = Path(arguments[0]).resolve() if arguments else ROOT / "dist/programa-limpieza-desinfeccion"
ZIP_OUTPUT = OUTPUT.parent / "extintores-scm-programa-limpieza-desinfeccion.zip"
TODAY = date.today().isoformat()

source = json.loads(SOURCE.read_text(encoding="utf-8"))
seed = json.loads(SEED.read_text(encoding="utf-8"))
company = seed["company"]
documents = source["documents"]
OUTPUT.mkdir(parents=True, exist_ok=True)


def rows_for(table: dict) -> list[list[str]]:
    headers = table["headers"]
    rows = [[str(cell) for cell in row] for row in table.get("rows", [])]
    rows.extend([[""] * len(headers) for _ in range(int(table.get("blank_rows", 0)))])
    if any(len(row) != len(headers) for row in rows):
        raise ValueError("Una fila no coincide con los encabezados de la plantilla")
    return rows


def shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    properties.append(fill)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build_program(document: dict) -> Path:
    word = Document()
    section = word.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    normal = word.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size, color in (("Title", 20, "123A5A"), ("Heading 1", 12, "176B87")):
        style = word.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
    header = section.header.paragraphs[0]
    header.text = f"{company['razonsocial']}  |  {document['code']}  |  Borrador"
    header.style = word.styles["Caption"]
    add_page_number(section.footer.paragraphs[0])

    title = word.add_paragraph(document["title"], style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = word.add_paragraph("Programa documental adaptable · SG-SST y saneamiento")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    metadata = word.add_table(rows=3, cols=4)
    metadata.style = "Table Grid"
    values = [
        ("Empresa", company["razonsocial"], "NIT", company["nit"]),
        ("Ciudad", company["ciudad"], "Código", document["code"]),
        ("Versión", "Borrador editable", "Fecha de generación", TODAY),
    ]
    for row, row_values in zip(metadata.rows, values):
        for index, value in enumerate(row_values):
            row.cells[index].text = value
            for run in row.cells[index].paragraphs[0].runs:
                run.font.size = Pt(8.5)
                run.font.bold = index % 2 == 0
            if index % 2 == 0:
                shade(row.cells[index], "EAF2F8")

    notice = word.add_table(rows=1, cols=1)
    notice.style = "Table Grid"
    notice.cell(0, 0).text = (
        "BORRADOR PARA ADAPTAR, APROBAR E IMPLEMENTAR. La creación o firma de este documento "
        "no demuestra capacitación, ejecución, verificación ni cierre del requerimiento IVC 2.3. "
        f"Referencia de origen: acta {seed['acta']['numero_acta']}, requerimiento 2.3."
    )
    shade(notice.cell(0, 0), "FFF4D6")
    for run in notice.cell(0, 0).paragraphs[0].runs:
        run.font.bold = True
        run.font.size = Pt(8.5)

    for section_data in document["sections"]:
        word.add_heading(section_data["title"], level=1)
        for text in section_data.get("paragraphs", []):
            word.add_paragraph(text)
        for text in section_data.get("unordered", []):
            word.add_paragraph(text, style="List Bullet")
        for text in section_data.get("ordered", []):
            word.add_paragraph(text, style="List Number")
        if "table" not in section_data:
            continue
        table_data = section_data["table"]
        headers = table_data["headers"]
        rows = rows_for(table_data)
        table = word.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = True
        for index, text in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = text
            shade(cell, "176B87")
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(7.5)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for values in rows:
            cells = table.add_row().cells
            for index, value in enumerate(values):
                cells[index].text = value if value else "\n"
                for paragraph in cells[index].paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.size = Pt(7.5)

    path = OUTPUT / "PR-LD-01_Programa_limpieza_desinfeccion_Extintores_SCM.docx"
    word.save(path)
    return path


def safe_sheet_name(value: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", value).strip()[:31] or "Formato"
    name = base
    suffix = 2
    while name in used:
        ending = f" {suffix}"
        name = base[: 31 - len(ending)] + ending
        suffix += 1
    used.add(name)
    return name


def build_form(document: dict, filename: str) -> Path:
    workbook = Workbook()
    control = workbook.active
    control.title = "Control"
    navy = "123A5A"
    teal = "176B87"
    pale = "EAF2F8"
    warning = "FFF4D6"
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    control_rows = [
        (document["title"], ""),
        ("Empresa", company["razonsocial"]),
        ("NIT", company["nit"]),
        ("Ciudad", company["ciudad"]),
        ("Código", document["code"]),
        ("Estado", "Borrador editable"),
        ("Fecha de generación", TODAY),
        ("Referencia", f"Acta {seed['acta']['numero_acta']} · requerimiento IVC 2.3"),
        ("Advertencia", "No acredita implementación. Diligenciar únicamente actividades y verificaciones reales; no antedatar registros."),
    ]
    for row in control_rows:
        control.append(row)
    control.merge_cells("A1:B1")
    control["A1"].font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    control["A1"].fill = PatternFill("solid", fgColor=navy)
    control["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    control.row_dimensions[1].height = 42
    for row in range(2, control.max_row + 1):
        control.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color=navy)
        control.cell(row, 1).fill = PatternFill("solid", fgColor=pale)
        for column in (1, 2):
            control.cell(row, column).border = border
            control.cell(row, column).alignment = Alignment(wrap_text=True, vertical="top")
    control.column_dimensions["A"].width = 24
    control.column_dimensions["B"].width = 95
    control.row_dimensions[9].height = 45
    control.freeze_panes = "A2"
    control.sheet_view.showGridLines = False
    # Las secciones sin tabla también deben quedar completas en el libro.
    for section_data in document["sections"]:
        if "table" in section_data:
            continue
        control.append([])
        title_row = control.max_row + 1
        control.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=2)
        control.cell(title_row, 1, section_data["title"])
        control.cell(title_row, 1).font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        control.cell(title_row, 1).fill = PatternFill("solid", fgColor=teal)
        texts = list(section_data.get("paragraphs", []))
        texts.extend(f"• {value}" for value in section_data.get("unordered", []))
        texts.extend(f"{index}. {value}" for index, value in enumerate(section_data.get("ordered", []), start=1))
        for text in texts:
            row = control.max_row + 1
            # openpyxl no considera celdas combinadas al calcular max_row en todos los casos.
            while control.cell(row, 1).value is not None:
                row += 1
            control.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            control.cell(row, 1, text)
            control.cell(row, 1).font = Font(name="Arial", size=9)
            control.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
            control.row_dimensions[row].height = max(28, 15 * (1 + len(text) // 150))

    used_names = {"Control"}
    for section_data in document["sections"]:
        if "table" not in section_data:
            continue
        table_data = section_data["table"]
        headers = table_data["headers"]
        rows = rows_for(table_data)
        sheet = workbook.create_sheet(safe_sheet_name(table_data.get("sheet", section_data["title"]), used_names))
        sheet.sheet_view.showGridLines = False
        width = len(headers)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        sheet.cell(1, 1, document["title"])
        sheet.cell(1, 1).font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        sheet.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
        sheet.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 38
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        sheet.cell(2, 1, f"{company['razonsocial']} · {document['code']} · Borrador · {TODAY}")
        sheet.cell(2, 1).font = Font(name="Arial", size=9, italic=True, color=navy)
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=width)
        sheet.cell(3, 1, section_data["title"])
        sheet.cell(3, 1).font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        sheet.cell(3, 1).fill = PatternFill("solid", fgColor=teal)
        current_row = 4
        for text in section_data.get("paragraphs", []):
            sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=width)
            cell = sheet.cell(current_row, 1, text)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = PatternFill("solid", fgColor=warning if current_row == 4 else "FFFFFF")
            cell.font = Font(name="Arial", size=9)
            sheet.row_dimensions[current_row].height = max(28, 15 * (1 + len(text) // 150))
            current_row += 1
        header_row = current_row
        for column, text in enumerate(headers, start=1):
            cell = sheet.cell(header_row, column, text)
            cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=teal)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = border
        sheet.row_dimensions[header_row].height = 42
        for row_values in rows:
            current_row += 1
            for column, value in enumerate(row_values, start=1):
                cell = sheet.cell(current_row, column, value)
                cell.font = Font(name="Arial", size=9)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border
            sheet.row_dimensions[current_row].height = 34 if any(row_values) else 42
        for column, text in enumerate(headers, start=1):
            content_lengths = [len(str(sheet.cell(row, column).value or "")) for row in range(header_row, current_row + 1)]
            sheet.column_dimensions[get_column_letter(column)].width = min(48, max(14, len(text) + 3, max(content_lengths, default=0) + 2))
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(width)}{current_row}"
        sheet.print_title_rows = f"1:{header_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_options.horizontalCentered = True
        sheet.oddFooter.center.text = f"{document['code']} · Borrador · Página &P de &N"

    path = OUTPUT / filename
    workbook.save(path)
    return path


def expected_strings(document: dict) -> list[str]:
    values = [document["title"]]
    for section in document["sections"]:
        values.append(section["title"])
        values.extend(section.get("paragraphs", []))
        values.extend(section.get("unordered", []))
        values.extend(section.get("ordered", []))
        if "table" in section:
            values.extend(section["table"]["headers"])
            values.extend(cell for row in rows_for(section["table"]) for cell in row if cell)
    return values


def verify_docx(path: Path, document: dict) -> None:
    word = Document(path)
    text = "\n".join(
        [paragraph.text for paragraph in word.paragraphs]
        + [cell.text for table in word.tables for row in table.rows for cell in row.cells]
    )
    missing = [value for value in expected_strings(document) if value not in text]
    if missing:
        raise ValueError(f"DOCX incompleto: {missing[0][:100]}")


def verify_xlsx(path: Path, document: dict) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    values = {
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }
    missing = [value for value in expected_strings(document) if value not in values]
    if missing:
        raise ValueError(f"XLSX incompleto: {missing[0][:100]}")
    if any(sheet.max_row > 1_048_576 or sheet.max_column > 16_384 for sheet in workbook.worksheets):
        raise ValueError("El libro excede los límites de Excel")


program_path = build_program(documents["program"])
execution_path = build_form(
    documents["execution"],
    "FT-LD-01_Programacion_registro_limpieza_desinfeccion_Extintores_SCM.xlsx",
)
verification_path = build_form(
    documents["verification"],
    "FT-LD-02_Verificacion_programa_limpieza_desinfeccion_Extintores_SCM.xlsx",
)
verify_docx(program_path, documents["program"])
verify_xlsx(execution_path, documents["execution"])
verify_xlsx(verification_path, documents["verification"])

readme = OUTPUT / "LEEME.txt"
readme.write_text(
    "EXTINTORES SCM · PROGRAMA DE LIMPIEZA Y DESINFECCIÓN\n"
    "===================================================\n\n"
    "Contenido:\n"
    "- PR-LD-01: programa formal editable en Word.\n"
    "- FT-LD-01: programación, ejecución y verificación en Excel.\n"
    "- FT-LD-02: lista de verificación y acciones en Excel.\n\n"
    "Uso requerido:\n"
    "1. Recorra las instalaciones y ajuste áreas, equipos, riesgos y frecuencias reales.\n"
    "2. Complete productos, diluciones, tiempos de contacto, EPP y SDS según fabricantes.\n"
    "3. Defina responsables, recursos, área de implementos, capacitación y cronograma.\n"
    "4. Revise y apruebe el programa. Los campos de aprobación inician en blanco.\n"
    "5. Ejecute y registre actividades reales; verifique soportes y gestione desviaciones.\n"
    "6. Adjunte en Normia SST las evidencias revisadas del requerimiento IVC 2.3.\n\n"
    "Los archivos se entregan como BORRADOR. No contienen ejecuciones, capacitaciones, "
    "aprobaciones, cierres ni evidencias inventadas. La existencia del kit no cierra el "
    "requerimiento 2.3 ni representa concepto favorable de la autoridad.\n\n"
    "Datos empresariales pendientes de validación documental: NIT 39626936-8, ciudad Soacha, "
    "clasificación CIIU y demás soportes del perfil. Las direcciones divergentes y los nombres "
    "de responsables deben conciliarse antes de aprobar.\n",
    encoding="utf-8",
)

files = [program_path, execution_path, verification_path, readme]
manifest = {
    "kit": source["kit"],
    "profile": "Extintores SCM",
    "status": "Borrador para adaptación, aprobación e implementación",
    "generated_on": TODAY,
    "requirement": "2.3",
    "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    "files": [
        {"name": path.name, "bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        for path in files
    ],
}
manifest_path = OUTPUT / "manifest.json"
manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
files.append(manifest_path)

with zipfile.ZipFile(ZIP_OUTPUT, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
    for path in files:
        archive.write(path, f"programa-limpieza-desinfeccion/{path.name}")
with zipfile.ZipFile(ZIP_OUTPUT) as archive:
    broken = archive.testzip()
    if broken:
        raise ValueError(f"Fallo CRC en {broken}")
    if any(path.suffix.lower() in {".exe", ".js", ".vbs", ".xlsm", ".docm"} for path in files):
        raise ValueError("El paquete contiene un formato ejecutable no permitido")

print(f"Kit verificado: {ZIP_OUTPUT}")
print(f"3 documentos editables + LEEME + manifest · {len(files)} archivos")
print(f"SHA-256: {hashlib.sha256(ZIP_OUTPUT.read_bytes()).hexdigest()}")
